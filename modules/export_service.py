import io
import logging
from typing import Dict, Any, List

# Docx imports
from docx import Document as DocxDocument
from docx.shared import Inches, Pt

# Pptx imports
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt

# ReportLab imports (PDF)
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

logger = logging.getLogger(__name__)

def export_to_docx(title: str, sections: Dict[str, str]) -> bytes:
    """
    Compiles academic sections into a Microsoft Word (.docx) document stream.
    
    Args:
        title: The primary title of the document.
        sections: Dictionary mapping section headings (e.g. 'INTRODUCTION') to text content.
        
    Returns:
        Bytes representation of the DOCX file.
    """
    try:
        logger.info(f"Compiling DOCX export for '{title}'.")
        doc = DocxDocument()
        
        # Style page title
        doc.add_heading(title, level=0)
        
        for heading, body in sections.items():
            if not body.strip():
                continue
            # Clean heading label
            clean_heading = heading.replace("#", "").replace("_", " ").strip().title()
            doc.add_heading(clean_heading, level=1)
            
            # Split paragraphs
            paragraphs = body.split("\n\n")
            for p in paragraphs:
                p_text = p.strip()
                if p_text.startswith("-") or p_text.startswith("*"):
                    # Render as list item
                    bullet_text = p_text.lstrip("-* ").strip()
                    doc.add_paragraph(bullet_text, style='List Bullet')
                elif p_text:
                    doc.add_paragraph(p_text)
                    
        stream = io.BytesIO()
        doc.save(stream)
        return stream.getvalue()
    except Exception as e:
        logger.error(f"Failed to generate DOCX file: {str(e)}", exc_info=True)
        raise RuntimeError(f"Error compiling DOCX: {str(e)}") from e

def export_to_pptx(slides_list: List[Dict[str, Any]]) -> bytes:
    """
    Compiles generated slide schemas into an actual Microsoft PowerPoint (.pptx) file.
    
    Args:
        slides_list: List of dicts representing slides. Each slide contains 'title', 
                     'bullets' (list of strings), and 'speaker_notes'.
                     
    Returns:
        Bytes representation of the PPTX file.
    """
    try:
        logger.info(f"Compiling PPTX export with {len(slides_list)} slides.")
        prs = Presentation()
        
        # Define layouts (Layout 5 is Title Only, Layout 6 is Blank, Layout 1 is Title + Content)
        title_layout = prs.slide_layouts[0]
        content_layout = prs.slide_layouts[1]
        
        # Add slide details
        for idx, slide_data in enumerate(slides_list):
            title_text = slide_data.get("title", f"Slide {idx + 1}").replace("SLIDE_TITLE:", "").strip()
            bullets = slide_data.get("bullets", [])
            notes = slide_data.get("speaker_notes", "").replace("SPEAKER_NOTES:", "").strip()
            
            # Use title slide layout for the first slide
            layout = title_layout if idx == 0 else content_layout
            slide = prs.slides.add_slide(layout)
            
            # Set slide Title
            title_placeholder = slide.shapes.title
            if title_placeholder:
                title_placeholder.text = title_text
                
            # Set Slide Content Bullets
            if bullets and len(slide.shapes) > 1:
                body_shape = slide.shapes[1]
                if body_shape.has_text_frame:
                    tf = body_shape.text_frame
                    tf.clear()
                    for bullet_idx, b_text in enumerate(bullets):
                        b_text = b_text.strip().lstrip("-* ").strip()
                        if not b_text:
                            continue
                        p = tf.add_paragraph() if bullet_idx > 0 else tf.paragraphs[0]
                        p.text = b_text
                        p.level = 0
            
            # Add Presenter Notes
            if notes:
                notes_slide = slide.notes_slide
                text_frame = notes_slide.notes_text_frame
                text_frame.text = notes
                
        stream = io.BytesIO()
        prs.save(stream)
        return stream.getvalue()
    except Exception as e:
        logger.error(f"Failed to generate PPTX file: {str(e)}", exc_info=True)
        raise RuntimeError(f"Error compiling PPTX: {str(e)}") from e

def export_to_pdf(title: str, sections: Dict[str, str]) -> bytes:
    """
    Compiles academic sections into a formatted PDF document.
    
    Args:
        title: Page title content.
        sections: Dictionary containing section headings and body paragraphs.
        
    Returns:
        Bytes representation of the PDF file.
    """
    try:
        logger.info(f"Compiling PDF report for '{title}'.")
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54
        )
        
        styles = getSampleStyleSheet()
        
        # Create premium typography styles
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=24,
            leading=28,
            textColor=colors.HexColor('#6366f1'),
            spaceAfter=20
        )
        
        h1_style = ParagraphStyle(
            'SectionHeader',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=20,
            textColor=colors.HexColor('#a855f7'),
            spaceBefore=14,
            spaceAfter=8,
            keepWithNext=True
        )
        
        body_style = ParagraphStyle(
            'BodyTextCustom',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#1e293b'),
            spaceAfter=10
        )
        
        bullet_style = ParagraphStyle(
            'BulletCustom',
            parent=body_style,
            leftIndent=20,
            bulletIndent=10,
            spaceAfter=6
        )
        
        story = []
        
        # Add Page Title
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 10))
        
        # Add Sections
        for heading, body in sections.items():
            if not body.strip():
                continue
                
            clean_heading = heading.replace("#", "").replace("_", " ").strip().title()
            story.append(Paragraph(clean_heading, h1_style))
            
            paragraphs = body.split("\n\n")
            for p in paragraphs:
                p_text = p.strip()
                if not p_text:
                    continue
                    
                if p_text.startswith("-") or p_text.startswith("*"):
                    bullet_text = p_text.lstrip("-* ").strip()
                    story.append(Paragraph(f"&bull; {bullet_text}", bullet_style))
                else:
                    story.append(Paragraph(p_text, body_style))
                    
        doc.build(story)
        pdf_data = buffer.getvalue()
        buffer.close()
        return pdf_data
    except Exception as e:
        logger.error(f"Failed to generate PDF report: {str(e)}", exc_info=True)
        raise RuntimeError(f"Error compiling PDF: {str(e)}") from e
